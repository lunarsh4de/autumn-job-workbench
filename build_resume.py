from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING


OUT = r"output\秋招初步简历_中文_图1样式.docx"

NAVY = "173B55"
TEAL = "2B7185"
MID_TEAL = "4B9BAB"
LIGHT_BLUE = "EAF5F8"
PALE_BLUE = "F5FAFC"
GRAY = "45515A"
LIGHT_GRAY = "D6E6EA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=5, inside=True):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right", "insideH", "insideV"]
    for name in names:
        tag = borders.find(qn(f"w:{name}"))
        if tag is None:
            tag = OxmlElement(f"w:{name}")
            borders.append(tag)
        tag.set(qn("w:val"), "single" if inside or name in ("top", "left", "bottom", "right") else "nil")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = borders.find(qn(f"w:{name}"))
        if tag is None:
            tag = OxmlElement(f"w:{name}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_run_font(run, name="Microsoft YaHei", size=9.0, bold=False, color=GRAY, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_border(p, side="bottom", color=MID_TEAL, size=8, space=3):
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    tag = p_bdr.find(qn(f"w:{side}"))
    if tag is None:
        tag = OxmlElement(f"w:{side}")
        p_bdr.append(tag)
    tag.set(qn("w:val"), "single")
    tag.set(qn("w:sz"), str(size))
    tag.set(qn("w:space"), str(space))
    tag.set(qn("w:color"), color)


def set_keep(p, with_next=False, together=False):
    p_pr = p._p.get_or_add_pPr()
    if with_next:
        el = OxmlElement("w:keepNext")
        p_pr.append(el)
    if together:
        el = OxmlElement("w:keepLines")
        p_pr.append(el)


def add_text(p, text, size=9.0, bold=False, color=GRAY, italic=False):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return r


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_paragraph_border(p, "left", color=TEAL, size=16, space=5)
    set_paragraph_border(p, "bottom", color=MID_TEAL, size=8, space=2)
    add_text(p, text, size=11.0, bold=True, color=NAVY)
    set_keep(p, with_next=True, together=True)
    return p


def add_label_line(doc, label, text, size=8.55):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, label, size=size, bold=True, color=NAVY)
    add_text(p, text, size=size, color=GRAY)
    set_keep(p, together=True)
    return p


def add_role_line(doc, left, date, size=9.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.04), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
    add_text(p, left, size=size, bold=True, color=NAVY)
    add_text(p, "\t" + date, size=size, bold=True, color=NAVY)
    set_keep(p, with_next=True, together=True)
    return p


def add_bullet(doc, label, text, size=8.55):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, "· ", size=size, bold=True, color=TEAL)
    if label:
        add_text(p, label, size=size, bold=True, color=NAVY)
    add_text(p, text, size=size, color=GRAY)
    set_keep(p, together=True)
    return p


def add_plain_line(doc, text, size=8.55, bold=False, color=GRAY, after=0.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, size=size, bold=bold, color=color)
    set_keep(p, together=True)
    return p


def set_page_border(section):
    sect_pr = section._sectPr
    borders = sect_pr.first_child_found_in("w:pgBorders")
    if borders is None:
        borders = OxmlElement("w:pgBorders")
        sect_pr.append(borders)
    borders.set(qn("w:offsetFrom"), "page")
    for name in ["top", "left", "bottom", "right"]:
        tag = borders.find(qn(f"w:{name}"))
        if tag is None:
            tag = OxmlElement(f"w:{name}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:space"), "6")
        tag.set(qn("w:color"), LIGHT_GRAY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.55)
    normal.font.color.rgb = RGBColor.from_string(GRAY)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def add_header(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.08)
    table.columns[1].width = Inches(6.15)
    remove_table_borders(table)
    left, right = table.rows[0].cells
    left.width = Inches(1.08)
    right.width = Inches(6.15)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(left, NAVY)
    set_cell_margins(left, top=100, start=90, bottom=100, end=90)
    set_cell_margins(right, top=30, start=170, bottom=30, end=40)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "YZ", size=23, bold=True, color=WHITE)
    p2 = left.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    add_text(p2, "工业设计", size=8.0, bold=True, color=WHITE)

    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "杨子辰", size=20, bold=True, color=NAVY)
    add_text(p, "  2027届硕士研究生", size=10.5, bold=True, color=TEAL)
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "求职方向：工业设计｜产品设计｜智能硬件", size=9.2, bold=True, color=GRAY)
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "电话：+852 5221 8490  ｜  邮箱：Lunarsh4de@gmail.com", size=8.7, color=GRAY)
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "核心优势｜工业设计、三维建模与渲染、3D 打印原型、嵌入式系统基础，具备从概念到样机验证的完整项目经验。", size=8.45, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    set_paragraph_border(p, "bottom", color=MID_TEAL, size=12, space=1)


def add_education(doc):
    add_section_heading(doc, "教育背景")
    add_role_line(doc, "香港大学｜创新设计与技术（IDT）理学硕士", "2026.09–至今")
    add_label_line(doc, "相关课程｜", "高级 CAD/CAM 与 AI 驱动制造系统；产品开发与应用编程；无人机设计、导航与控制。")
    add_role_line(doc, "宁波工程学院｜工业设计 工学学士", "2020.09–2024.06")
    add_label_line(doc, "专业基础｜", "工业设计、产品开发、三维建模、视觉表达、交互设计、原型制作与基础嵌入式开发。")


def add_work(doc):
    add_section_heading(doc, "工作经历")
    add_role_line(doc, "宁波智迪自动化科技有限公司｜工业设计师", "2025.09–2026.06")
    add_bullet(doc, "产品外壳设计｜", "为电机驱动器设计外壳，与工程团队协同确定外观与结构，制作 3D 打印样机并完成设计验证。")
    add_bullet(doc, "创意产品设计｜", "根据客户需求开发展会定制宣传礼品，负责三维建模、打样与 3D 打印。")
    add_role_line(doc, "中国兵器工业集团航联科技有限公司｜设计助理", "2024.06–2024.09")
    add_bullet(doc, "结构设计｜", "参与天线防护罩设计，提出结构改进方案，完成 3 轮整体结构迭代、建模与验证。")
    add_bullet(doc, "设计支持｜", "协助优化设备外壳外观并制作效果图，用于设计评审与汇报；形成兼顾防护功能与结构可行性的方案。")
    add_role_line(doc, "宁波一间设计研究院｜助理技术员", "2022.06–2022.09")
    add_bullet(doc, "项目产出｜", "负责多个项目的概念开发、建模、渲染及版式输出；代表项目包括管道清洁车、时光烘茶杯、面向情绪需求的造口袋设计研究。")


def add_projects(doc):
    add_section_heading(doc, "项目经历")
    add_role_line(doc, "智能下水道清洁机器人｜核心成员", "2023.09–2024.04")
    add_bullet(doc, "项目实现｜", "使用 UG/NX 完成三维建模与 KeyShot 渲染；编写 Arduino 初版控制程序，3D 打印车体并开展性能评估；利用 ComfyUI 模拟城市街景测试场景。")
    add_role_line(doc, "国家级大学生创新训练计划｜第二负责人", "2022.09–2023.06")
    add_bullet(doc, "方案开发｜", "开发基于液压循环技术的分布式投饵系统，参与概念确定、产品设计、测试与模具打样全流程。")
    add_bullet(doc, "技术成果｜", "通过电动三通阀实现分仓投饵自动控制，形成低成本、高效率的螃蟹投喂方案。")
    add_role_line(doc, "专利项目：自动投饵器｜设计负责人", "2022.01–2022.12")
    add_bullet(doc, "研发工作｜", "开展行业调研和用户需求建模，协助完成从市场分析到样机测试的研发流程，起草振动投料机构专利申请。")
    add_role_line(doc, "智能消毒车｜团队负责人", "2021.09–2022.06")
    add_bullet(doc, "团队与工程｜", "组织机械、电子与工业设计跨学科团队完成机器人开发；集成视觉模块实现动态路径识别；使用 Arduino 编程、3D 打印制作部件并装配可运行样机。")


def add_skills(doc):
    add_section_heading(doc, "岗位匹配与专业技能")
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.42, 2.42, 2.42]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    set_table_borders(table, color=LIGHT_GRAY, size=5, inside=True)
    headers = ["三维设计与视觉", "原型与嵌入式", "工具与工程协作"]
    values = [
        "Siemens NX（UG）、Rhino、KeyShot；Adobe Photoshop、Illustrator、Figma",
        "Arduino、3D 打印、机械装配；电机/传感器基础集成、基础控制逻辑",
        "Python、ComfyUI；产品概念、结构迭代、样机验证、效果图与版式输出",
    ]
    for i in range(3):
        h = table.cell(0, i)
        v = table.cell(1, i)
        set_cell_shading(h, LIGHT_BLUE)
        set_cell_shading(v, PALE_BLUE)
        set_cell_margins(h, top=60, start=90, bottom=45, end=90)
        set_cell_margins(v, top=45, start=90, bottom=55, end=90)
        hp = h.paragraphs[0]
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.line_spacing = 1.0
        add_text(hp, headers[i], size=8.35, bold=True, color=NAVY)
        vp = v.paragraphs[0]
        vp.paragraph_format.space_after = Pt(0)
        vp.paragraph_format.line_spacing = 1.0
        add_text(vp, values[i], size=7.85, color=GRAY)


def add_patents_honors(doc):
    add_section_heading(doc, "专利与荣誉")
    add_plain_line(doc, "专利｜共同发明人：自动投饵器（ZL 2022 1 0406391.1），2022；发明人：物理教学辅助储物装置（ZL 2019 2 1194467.9），2020。", size=8.35)
    add_plain_line(doc, "荣誉｜国际级：第4届“金紫荆杯”中国—东盟工业设计大赛入围奖（2022）；国家级：国家级大学生创新训练计划（2023）；省级：“互联网+”创新创业大赛铜奖（2022）；市级：宁波市“37℃爱”设计大赛一等奖（2022）；校级：三等奖学金（2023）、优秀团员（2020/2021/2022）。", size=8.35)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.36)
    section.bottom_margin = Inches(0.36)
    section.left_margin = Inches(0.47)
    section.right_margin = Inches(0.47)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)
    set_page_border(section)
    configure_styles(doc)

    add_header(doc)
    add_education(doc)
    add_work(doc)
    add_projects(doc)
    add_skills(doc)
    add_patents_honors(doc)

    # Keep the document intentionally compact and visually dense like the reference.
    core = doc.core_properties
    core.title = "杨子辰｜中文简历"
    core.subject = "工业设计 / 产品设计 / 智能硬件"
    core.author = "杨子辰"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
