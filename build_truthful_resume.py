from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\81489\Documents\求职\output\YANG Zichen CV_工业设计优化版.docx")

NAVY = "193B53"
BLUE = "2D667C"
TEAL = "2A7D78"
TEXT = "27323A"
MUTED = "5E6A71"
RULE = "B9CED6"
PALE = "EAF3F5"


def set_run_font(run, size=9.2, bold=False, color=TEXT, italic=False):
    run.font.name = "Aptos"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_run(paragraph, text, size=9.2, bold=False, color=TEXT, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def set_keep(paragraph, next_paragraph=False, together=True):
    ppr = paragraph._p.get_or_add_pPr()
    if together and ppr.find(qn("w:keepLines")) is None:
        ppr.append(OxmlElement("w:keepLines"))
    if next_paragraph and ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def set_bottom_border(paragraph, color=RULE, size=7, space=2):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def set_left_border(paragraph, color=TEAL, size=14, space=7):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    pbdr.append(left)


def add_bullet_numbering(document):
    numbering = document.part.numbering_part.element
    used_abstract = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    used_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(used_abstract, default=0) + 1
    num_id = max(used_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•")
    level.append(text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    ind.set(qn("w:hanging"), "180")
    ppr.append(ind)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    rpr.append(color)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.insert(0, numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(11.5)
    section.bottom_margin = Mm(10.5)
    section.left_margin = Mm(14.5)
    section.right_margin = Mm(14.5)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(1.4)
    normal.paragraph_format.line_spacing = 1.03

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor.from_string(NAVY)

    core = document.core_properties
    core.title = "YANG Zichen 杨梓琛｜工业设计中文简历"
    core.subject = "2027届｜工业设计、产品设计、三维设计与快速原型"
    core.author = "杨梓琛"
    core.keywords = "工业设计, 产品设计, 三维建模, NX, Rhino, KeyShot, 3D打印, Arduino, 样机验证"


def add_header(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_run(p, "YANG Zichen  杨梓琛", size=20, bold=True, color=NAVY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1.5)
    add_run(p, "2027届硕士研究生  |  中共党员  |  江苏泰州", size=9.3, color=MUTED)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    add_run(p, "15861095585  |  Lunarsh4de@gmail.com", size=9.3, color=MUTED)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "工业设计  |  产品设计  |  三维设计与快速原型", size=10.2, bold=True, color=TEAL)
    set_bottom_border(p, color=TEAL, size=10, space=3)


def add_profile(document):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Mm(2)
    p.paragraph_format.right_indent = Mm(1)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    set_left_border(p)
    add_run(
        p,
        "工业设计本科，香港大学创新设计与技术硕士在读。具备 NX/Rhino 三维建模、KeyShot 渲染和 3D 打印原型经验；完成过电机驱动器外壳、非标支架与定制产品设计，参与结构迭代、样机验证及电机/电缸调试，可支持从概念表达、三维建模到实体验证的设计开发环节。",
        size=9.2,
        color=TEXT,
    )


def add_section(document, title):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(3.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, title, size=11.2, bold=True, color=NAVY)
    set_bottom_border(p, color=RULE, size=6, space=2)
    set_keep(p, next_paragraph=True)


def add_entry(document, left, date, subtitle=None):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(1.4)
    p.paragraph_format.space_after = Pt(0.6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, left, size=9.5, bold=True, color=NAVY)
    add_run(p, "\t" + date, size=9.1, bold=True, color=BLUE)
    set_keep(p, next_paragraph=subtitle is not None)
    if subtitle:
        s = document.add_paragraph()
        s.paragraph_format.space_after = Pt(0.7)
        add_run(s, subtitle, size=8.9, color=MUTED, italic=True)
        set_keep(s, next_paragraph=True)


def add_bullet(document, num_id, text, lead=None):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.9)
    p.paragraph_format.line_spacing = 1.03
    apply_bullet(p, num_id)
    if lead:
        add_run(p, lead, size=9.15, bold=True, color=BLUE)
    add_run(p, text, size=9.15, color=TEXT)
    set_keep(p)


def add_label_line(document, label, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.2)
    p.paragraph_format.line_spacing = 1.02
    add_run(p, label, size=9.1, bold=True, color=BLUE)
    add_run(p, text, size=9.1, color=TEXT)
    set_keep(p)


def build_resume():
    document = Document()
    configure_document(document)
    bullet_num = add_bullet_numbering(document)

    add_header(document)
    add_profile(document)

    add_section(document, "教育背景")
    add_entry(document, "香港大学｜工程学院｜创新设计与技术 硕士", "2026.09-至今（2027届）")
    add_label_line(document, "相关课程｜", "高级 CAD/CAM 与 AI 驱动制造系统、产品开发与应用编程、无人机设计/导航与控制")
    add_entry(document, "浙江大学宁波理工学院｜工业设计 工学学士", "2020.09-2024.06")

    add_section(document, "专业技能")
    add_label_line(document, "三维与表达｜", "NX（UG）、Rhino、KeyShot；Photoshop、Illustrator、Figma")
    add_label_line(document, "原型与技术｜", "3D 打印、基础机械装配、Arduino 基础控制；Python、C++ 基础")
    add_label_line(document, "AIGC 工具｜", "ComfyUI 用于概念探索与场景视觉化；Codex 基础使用")

    add_section(document, "工作经历")
    add_entry(document, "宁波致迪自动化科技有限公司｜工业设计师", "2025.09-2026.06")
    add_bullet(document, bullet_num, "围绕外观、基础结构与散热约束完成电机驱动器外壳方案，配合研发确定结构，制作 3D 打印样机，参与设计验证与初步打样。", "外壳与样机｜")
    add_bullet(document, bullet_num, "为 2 款电机和 1 个电缸设计非标支架，完成三维建模与 3D 打印；设计电机支架-铝型材快速测试系统，配合研发进行参数调试。", "非标支架｜")
    add_bullet(document, bullet_num, "协助搭建电机、传感器与传送带组成的飞剪功能 Demo；根据客户需求完成展会定制礼品的建模、打样与 3D 打印。", "设计支持｜")

    add_entry(document, "中国兵器工业集团航联科技有限公司｜设计助理", "2023.06-2023.09")
    add_bullet(document, bullet_num, "参与天线防护罩结构设计，提出结构改进方案并完成 2 轮整体迭代、三维建模与验证；协助优化外观并制作评审效果图，支持方案评审与汇报。")

    add_entry(document, "宁波一尖设计学院｜助理工业设计师", "2022.06-2022.09")
    add_bullet(document, bullet_num, "参与多个项目的概念开发、建模、渲染与版式输出，包括管道清洁车、便携式烘茶杯及面向情绪需求的造口袋设计研究。")

    add_section(document, "项目经历")
    add_entry(document, "智能下水道清洁机器人｜核心成员", "2023.09-2024.04")
    add_bullet(document, bullet_num, "使用 UG/NX 完成三维建模、KeyShot 渲染与 3D 打印车体；编写 Arduino 初版控制程序并参与原型性能评估，使用 ComfyUI 完成城市街景场景视觉化。")

    add_entry(document, "青蟹养殖系统｜第二负责人｜国家级大学生创新训练计划", "2022.09-2023.06")
    add_bullet(document, bullet_num, "参与基于液压循环技术的分布式投饵系统开发，覆盖概念确定、产品设计、测试与模具打样；通过电动三通阀实现分仓投饵自动控制。")

    add_entry(document, "自动投饵器｜设计负责人｜发明专利项目", "2022.01-2022.12")
    add_bullet(document, bullet_num, "开展行业调研和用户需求建模，协助推进从市场分析到样机测试的研发流程，并起草振动投料机构专利申请。")

    add_entry(document, "智能消毒车｜团队负责人", "2021.09-2022.06")
    add_bullet(document, bullet_num, "组织机械、电子与工业设计跨学科团队完成机器人开发；参与视觉模块集成与动态路径识别，使用 Arduino 编程、3D 打印制作部件并装配可运行样机。")

    add_section(document, "专利与奖项")
    add_label_line(document, "专利｜", "发明专利“一款自动投饵机”（ZL 2022 1 0406391.1，2022）；实用新型专利“新型便于存放的物理实用教具”（第一发明人，ZL 2019 2 1194467.9，2020）")
    add_label_line(document, "奖项｜", "中国-东盟工业设计大赛“金紫荆杯”入围奖（2023）；全国大学生“互联网+”创新大赛浙江省铜奖（2022）；全国大学生英语翻译大赛省级一等奖（2022）")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
